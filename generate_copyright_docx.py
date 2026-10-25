import os
import sys
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_copyright_document(output_path):
    doc = docx.Document()
    
    # Page setup: Standard A4 with 0.75 in margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
        # Configure Header & Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("SKILL CONNECT — COMPLETE SOURCE CODE (COPYRIGHT REGISTRATION)")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = RGBColor(140, 140, 140)
        
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Department of Computer Science and Engineering · Sanjivani University · Copyright © 2026")
        frun.font.name = "Arial"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(140, 140, 140)

    # Palette
    COLOR_PRIMARY = RGBColor(30, 41, 59)     # Navy slate
    COLOR_ACCENT = RGBColor(98, 84, 220)     # Indigo accent
    COLOR_MUTED = RGBColor(100, 116, 139)    # Gray
    COLOR_DARK = RGBColor(15, 23, 42)        # Ink

    # =========================================================================
    # COVER PAGE
    # =========================================================================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("SANJIVANI UNIVERSITY\n")
    r_inst.font.name = "Arial"
    r_inst.font.size = Pt(14)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    r_dept = p_inst.add_run("School of Engineering and Technology\nDepartment of Computer Science and Engineering\n")
    r_dept.font.name = "Arial"
    r_dept.font.size = Pt(10)
    r_dept.font.color.rgb = COLOR_MUTED

    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_badge.paragraph_format.space_before = Pt(18)
    p_badge.paragraph_format.space_after = Pt(18)
    r_badge = p_badge.add_run("APPLICATION FOR COPYRIGHT REGISTRATION — FORM XIV\n[EXTRACT / COMPLETE REPOSITORY OF COMPUTER SOFTWARE SOURCE CODE]")
    r_badge.font.name = "Arial"
    r_badge.font.size = Pt(10)
    r_badge.font.bold = True
    r_badge.font.color.rgb = COLOR_ACCENT

    # Title Box
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(4)
    r_title = p_title.add_run("SKILL CONNECT\n")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    r_sub = p_title.add_run("A Campus Peer-to-Peer Skill Exchange and Collaborative Learning Platform\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_ACCENT

    # Authors Table
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    p_auth_head = doc.add_paragraph()
    p_auth_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ah = p_auth_head.add_run("AUTHORS / INVENTORS / APPLICANTS")
    r_ah.font.name = "Arial"
    r_ah.font.size = Pt(10)
    r_ah.font.bold = True
    r_ah.font.color.rgb = COLOR_PRIMARY

    table_auth = doc.add_table(rows=5, cols=4)
    table_auth.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_auth.autofit = False

    headers = ["Sr. No.", "Full Name of Author", "PRN / Student ID", "Role / Contribution"]
    widths = [Inches(0.8), Inches(2.5), Inches(1.8), Inches(1.6)]
    
    hdr_cells = table_auth.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].font.bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.size = Pt(9)
        hdr_cells[i].paragraphs[0].runs[0].font.name = "Arial"
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        hdr_cells[i].width = widths[i]

    authors_data = [
        ("1", "Mahesh Bangar", "2125UCEM1012", "Lead Architect & UI/UX"),
        ("2", "Saksham Patil", "2125UCEM1034", "Frontend & Auth Flow"),
        ("3", "Saisha Asane", "2125UCEF1015", "Matching & Data Engine"),
        ("4", "Chandana Bothara", "2125UCEF1051", "Communities & Sessions")
    ]

    for row_idx, data in enumerate(authors_data, start=1):
        row_cells = table_auth.rows[row_idx].cells
        for col_idx, text_val in enumerate(data):
            row_cells[col_idx].text = text_val
            p = row_cells[col_idx].paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.name = "Arial"
                p.runs[0].font.size = Pt(8.5)
                if col_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bg = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
            set_cell_background(row_cells[col_idx], bg)
            set_cell_margins(row_cells[col_idx], top=60, bottom=60, left=100, right=100)
            row_cells[col_idx].width = widths[col_idx]

    # Faculty Guidance Section
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    p_fac = doc.add_paragraph()
    p_fac.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_fac_title = p_fac.add_run("ACADEMIC SUPERVISION\n")
    r_fac_title.font.name = "Arial"
    r_fac_title.font.size = Pt(9.5)
    r_fac_title.font.bold = True
    r_fac_title.font.color.rgb = COLOR_PRIMARY

    r_fac = p_fac.add_run(
        "Project Guide: Prof. Jyotimoyee Kalita\n"
        "Head of Department: Dr. Mahendra Gawali\n"
        "Academic Year: 2026–2027 · Seminar-I Project"
    )
    r_fac.font.name = "Arial"
    r_fac.font.size = Pt(8.5)
    r_fac.font.color.rgb = COLOR_MUTED

    # Declaration of Originality
    p_dec = doc.add_paragraph()
    p_dec.paragraph_format.space_before = Pt(14)
    p_dec.paragraph_format.space_after = Pt(6)
    r_dh = p_dec.add_run("DECLARATION OF ORIGINAL WORK & COPYRIGHT STATEMENT\n")
    r_dh.font.name = "Arial"
    r_dh.font.size = Pt(9)
    r_dh.font.bold = True
    r_dh.font.color.rgb = COLOR_PRIMARY

    r_dt = p_dec.add_run(
        "We hereby declare that this software code titled \"SKILL CONNECT\" is an original computer software "
        "created and authored by us. The work has not been copied or adapted from any third-party copyrighted work. "
        "This document contains the verbatim source code files of the platform submitted for official copyright registration "
        "under the Copyright Act, 1957 (Government of India)."
    )
    r_dt.font.name = "Arial"
    r_dt.font.size = Pt(8)
    r_dt.font.color.rgb = COLOR_MUTED

    # Signatures
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(18)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sig = p_sig.add_run(
        "____________________          ____________________          ____________________          ____________________\n"
        "   Mahesh Bangar                  Saksham Patil                  Saisha Asane                Chandana Bothara\n"
        " (2125UCEM1012)                  (2125UCEM1034)                 (2125UCEF1015)                (2125UCEF1051)"
    )
    r_sig.font.name = "Arial"
    r_sig.font.size = Pt(7.5)
    r_sig.font.color.rgb = COLOR_PRIMARY

    doc.add_page_break()

    # =========================================================================
    # SECTION 1: SOFTWARE SPECIFICATIONS & INDEX
    # =========================================================================
    h1 = doc.add_heading("1. Software Specification & Project Overview", level=1)
    h1.runs[0].font.name = "Arial"
    h1.runs[0].font.color.rgb = COLOR_PRIMARY
    h1.runs[0].font.size = Pt(15)

    p_abs = doc.add_paragraph()
    p_abs.paragraph_format.space_after = Pt(8)
    r_ab = p_abs.add_run(
        "Skill Connect is a campus-dedicated peer-to-peer skill exchange platform designed to enable reciprocal, "
        "non-monetary learning among university students. The application operates as a Single-Page Application (SPA) "
        "engineered with React 19, modern CSS3 token architecture, and a custom zero-dependency client-side routing engine. "
        "The software incorporates 11 integrated modules: Institutional Login, 4-Step Profile Onboarding Wizard, Metric-Aggregating "
        "Dashboard, Algorithmic Reciprocal Matching, Campus Circles (Communities), Workshop Session Scheduling, August "
        "Interactive Calendar, Live Peer Messaging, Dynamic Campus Demand Polls, Keyboard-Driven Command Palette (Ctrl+K Search), "
        "and Comprehensive Profile Customization."
    )
    r_ab.font.name = "Arial"
    r_ab.font.size = Pt(9.5)
    r_ab.font.color.rgb = COLOR_DARK

    # Technical Characteristics Table
    p_th = doc.add_paragraph()
    p_th.paragraph_format.space_before = Pt(8)
    r_th = p_th.add_run("TECHNICAL CHARACTERISTICS OF THE COMPUTER SOFTWARE")
    r_th.font.name = "Arial"
    r_th.font.size = Pt(9.5)
    r_th.font.bold = True
    r_th.font.color.rgb = COLOR_PRIMARY

    table_tech = doc.add_table(rows=7, cols=2)
    table_tech.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_data = [
        ("Title of the Software Work", "SKILL CONNECT (Campus Skill Exchange Platform)"),
        ("Category of Work", "Computer Software / Literary Work (Under Copyright Act, 1957)"),
        ("Programming / Scripting Languages", "JavaScript (ES2022 / JSX), HTML5, CSS3"),
        ("Frontend Framework & Build System", "React 19.x, Vite 8.2.x, Lucide-React SVG Engine"),
        ("State Architecture & Persistence", "Reactive React Hooks with Persistent Browser LocalStorage"),
        ("Operating Environment / Target Platforms", "Cross-Platform: Windows, Linux, macOS, Android, iOS (Web-based)"),
        ("Source Code Repository Layout", "8 Dedicated Modules / Files (~806 Total Lines of Clean Source Code)")
    ]
    for idx, (label, val) in enumerate(tech_data):
        r = table_tech.rows[idx]
        r.cells[0].text = label
        r.cells[0].paragraphs[0].runs[0].font.bold = True
        r.cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
        r.cells[0].paragraphs[0].runs[0].font.name = "Arial"
        r.cells[0].width = Inches(2.6)
        set_cell_background(r.cells[0], "F1F5F9")
        set_cell_margins(r.cells[0], top=50, bottom=50, left=80, right=80)

        r.cells[1].text = val
        r.cells[1].paragraphs[0].runs[0].font.size = Pt(8.5)
        r.cells[1].paragraphs[0].runs[0].font.name = "Arial"
        r.cells[1].width = Inches(4.1)
        set_cell_background(r.cells[1], "FFFFFF")
        set_cell_margins(r.cells[1], top=50, bottom=50, left=80, right=80)

    # Index of Source Files Table
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    p_ih = doc.add_paragraph()
    r_ih = p_ih.add_run("INDEX OF REGISTERED SOURCE CODE FILES")
    r_ih.font.name = "Arial"
    r_ih.font.size = Pt(9.5)
    r_ih.font.bold = True
    r_ih.font.color.rgb = COLOR_PRIMARY

    files_info = [
        ("1", "package.json", "JSON", "407 B", "16", "Project dependencies, scripts, and build configuration"),
        ("2", "index.html", "HTML5", "73 B", "1", "Main application HTML entry point and viewport setup"),
        ("3", "src/main.jsx", "JavaScript (React)", "284 B", "8", "Client entry bootstrapper rendering App in Router"),
        ("4", "src/App.jsx", "JavaScript (React JSX)", "48.5 KB", "628", "Core application components, 11 modules, state machine"),
        ("5", "src/router.jsx", "JavaScript", "2.99 KB", "68", "Zero-dependency client-side SPA routing system"),
        ("6", "src/data.js", "JavaScript", "5.57 KB", "32", "Campus entities: taxonomy, peers, communities, sessions"),
        ("7", "src/styles.css", "CSS3", "35.3 KB", "34", "Design tokens, color themes, responsive layouts, components"),
        ("8", "src/refinement.css", "CSS3", "9.79 KB", "19", "Visual polish, lighting, shadow layers, theme refinement")
    ]

    table_index = doc.add_table(rows=len(files_info) + 1, cols=6)
    table_index.alignment = WD_TABLE_ALIGNMENT.CENTER
    idx_headers = ["No.", "File Relative Path", "Language", "File Size", "Lines", "Functional Description"]
    idx_widths = [Inches(0.4), Inches(1.8), Inches(1.2), Inches(0.8), Inches(0.6), Inches(2.0)]
    
    for i, h_text in enumerate(idx_headers):
        cell = table_index.rows[0].cells[i]
        cell.text = h_text
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)
        cell.paragraphs[0].runs[0].font.name = "Arial"
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(cell, "1E293B")
        set_cell_margins(cell, top=60, bottom=60, left=60, right=60)
        cell.width = idx_widths[i]

    for r_idx, f_data in enumerate(files_info, start=1):
        row = table_index.rows[r_idx]
        for c_idx, val in enumerate(f_data):
            c = row.cells[c_idx]
            c.text = val
            p = c.paragraphs[0]
            if len(p.runs) > 0:
                p.runs[0].font.name = "Arial"
                p.runs[0].font.size = Pt(8)
                if c_idx in [0, 3, 4]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
            set_cell_background(c, bg)
            set_cell_margins(c, top=40, bottom=40, left=60, right=60)
            c.width = idx_widths[c_idx]

    doc.add_page_break()

    # =========================================================================
    # SECTION 2: VERBATIM COMPLETE SOURCE CODE
    # =========================================================================
    h_code = doc.add_heading("2. Complete Source Code Repository", level=1)
    h_code.runs[0].font.name = "Arial"
    h_code.runs[0].font.color.rgb = COLOR_PRIMARY
    h_code.runs[0].font.size = Pt(15)

    p_c_desc = doc.add_paragraph()
    p_c_desc.paragraph_format.space_after = Pt(12)
    r_cd = p_c_desc.add_run(
        "The following sections provide the complete, unedited source code for all 8 files comprising the Skill Connect software platform. "
        "Line numbers are sequentially prefixed (e.g., 0001 | ) for exact reference during copyright examination."
    )
    r_cd.font.name = "Arial"
    r_cd.font.size = Pt(9)
    r_cd.font.color.rgb = COLOR_MUTED

    file_paths = [
        ("package.json", "package.json", "JSON Configuration & Dependencies"),
        ("index.html", "index.html", "HTML5 Entry Point"),
        ("src/main.jsx", "src/main.jsx", "Application Bootstrapper & Mount"),
        ("src/App.jsx", "src/App.jsx", "Core Application Architecture, Navigation & Modules"),
        ("src/router.jsx", "src/router.jsx", "Custom Client Router Engine"),
        ("src/data.js", "src/data.js", "Taxonomy & Campus Network Dataset"),
        ("src/styles.css", "src/styles.css", "Design Tokens, Themes & Responsive Stylesheet"),
        ("src/refinement.css", "src/refinement.css", "Visual Polish & Theme Refinements")
    ]

    for file_idx, (rel_path, display_name, file_desc) in enumerate(file_paths, start=1):
        if not os.path.exists(rel_path):
            print(f"Warning: File {rel_path} not found!")
            continue

        with open(rel_path, "r", encoding="utf-8") as fp:
            code_lines = fp.readlines()

        # Section Heading for File
        h_file = doc.add_heading(f"2.{file_idx} File: {display_name}", level=2)
        h_file.runs[0].font.name = "Arial"
        h_file.runs[0].font.color.rgb = COLOR_ACCENT
        h_file.runs[0].font.size = Pt(12)
        h_file.paragraph_format.space_before = Pt(12)
        h_file.paragraph_format.space_after = Pt(4)

        # File Meta Table
        table_fm = doc.add_table(rows=2, cols=3)
        table_fm.alignment = WD_TABLE_ALIGNMENT.CENTER
        fm_widths = [Inches(2.5), Inches(2.2), Inches(2.0)]
        
        row0 = table_fm.rows[0]
        row0.cells[0].text = f"Relative Path: {rel_path}"
        row0.cells[1].text = f"File Size: {os.path.getsize(rel_path):,} Bytes"
        row0.cells[2].text = f"Total Lines: {len(code_lines):,} Lines"

        row1 = table_fm.rows[1]
        row1.cells[0].text = f"Module Role: {file_desc}"
        row1.cells[1].text = f"Format: UTF-8 Plaintext"
        row1.cells[2].text = f"Author: Mahesh Bangar et al."

        for r in table_fm.rows:
            for c_i, cell in enumerate(r.cells):
                cell.width = fm_widths[c_i]
                set_cell_background(cell, "F1F5F9")
                set_cell_margins(cell, top=30, bottom=30, left=50, right=50)
                p = cell.paragraphs[0]
                if len(p.runs) > 0:
                    p.runs[0].font.name = "Arial"
                    p.runs[0].font.size = Pt(7.5)
                    p.runs[0].font.color.rgb = COLOR_PRIMARY

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

        # Code Block
        # We write code lines in chunks to prevent creating thousands of paragraphs
        # A single paragraph with newlines or small chunks maintains fast performance
        CHUNK_SIZE = 100
        for chunk_start in range(0, len(code_lines), CHUNK_SIZE):
            chunk = code_lines[chunk_start:chunk_start + CHUNK_SIZE]
            chunk_text = ""
            for line_no, line_content in enumerate(chunk, start=chunk_start + 1):
                clean_content = line_content.rstrip("\r\n")
                chunk_text += f"{line_no:04d} | {clean_content}\n"
            
            p_code = doc.add_paragraph()
            p_code.paragraph_format.space_before = Pt(0)
            p_code.paragraph_format.space_after = Pt(0)
            p_code.paragraph_format.line_spacing = Pt(9.5)
            r_code = p_code.add_run(chunk_text.rstrip("\n"))
            r_code.font.name = "Consolas"
            r_code.font.size = Pt(7.5)
            r_code.font.color.rgb = COLOR_DARK

        doc.add_page_break()

    # Save document
    doc.save(output_path)
    print(f"Successfully generated: {output_path}")

if __name__ == "__main__":
    out_file = "Skill_Connect_Source_Code_Copyright.docx"
    create_copyright_document(out_file)
